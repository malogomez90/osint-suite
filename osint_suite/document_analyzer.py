"""
Analizador de Documentos - Extracción de metadatos de PDF y Office
Excluye: Explotación de vulnerabilidades, extracción de macros maliciosas
"""

import os
import re
import json
from typing import Dict, List, Optional, Any
from datetime import datetime
from .utils import (
    Colors, print_banner, print_success, print_error, 
    print_warning, print_info, save_to_json
)


class DocumentAnalyzer:
    """Analizador de metadatos de documentos PDF y Office"""
    
    def __init__(self):
        self.results = {}
        self.supported_formats = {
            '.pdf': 'PDF',
            '.doc': 'Word (Legacy)',
            '.docx': 'Word (Modern)',
            '.xls': 'Excel (Legacy)',
            '.xlsx': 'Excel (Modern)',
            '.ppt': 'PowerPoint (Legacy)',
            '.pptx': 'PowerPoint (Modern)',
            '.odt': 'OpenDocument Text',
            '.ods': 'OpenDocument Spreadsheet',
            '.odp': 'OpenDocument Presentation',
            '.rtf': 'Rich Text Format'
        }
    
    def analyze_pdf(self, file_path: str) -> Dict:
        """
        Analiza metadatos de un archivo PDF
        
        Args:
            file_path: Ruta al archivo PDF
        
        Returns:
            Dict con metadatos del PDF
        """
        print_info(f"Analizando PDF: {os.path.basename(file_path)}")
        
        result = {
            'file_path': file_path,
            'file_name': os.path.basename(file_path),
            'file_type': 'PDF',
            'metadata': {},
            'xmp_metadata': {},
            'error': None,
            'security_info': {},
            'structure_info': {}
        }
        
        try:
            # Intentar usar PyPDF2
            try:
                from PyPDF2 import PdfReader
                
                reader = PdfReader(file_path)
                
                # Metadatos básicos
                if reader.metadata:
                    meta = reader.metadata
                    result['metadata'] = {
                        'title': meta.get('/Title', ''),
                        'author': meta.get('/Author', ''),
                        'subject': meta.get('/Subject', ''),
                        'creator': meta.get('/Creator', ''),
                        'producer': meta.get('/Producer', ''),
                        'creation_date': meta.get('/CreationDate', ''),
                        'modification_date': meta.get('/ModDate', ''),
                    }
                
                # Información del documento
                result['structure_info'] = {
                    'num_pages': len(reader.pages),
                    'is_encrypted': reader.is_encrypted,
                    'form_fields': len(reader.get_fields()) if reader.get_fields() else 0
                }
                
                # Intentar extraer texto de la primera página para análisis
                if len(reader.pages) > 0:
                    try:
                        first_page_text = reader.pages[0].extract_text()[:500]
                        result['structure_info']['first_page_preview'] = first_page_text
                    except:
                        pass
                
                print_success(f"PDF analizado: {result['structure_info']['num_pages']} páginas")
                
            except ImportError:
                print_warning("PyPDF2 no instalado, usando análisis básico")
                result['error'] = 'PyPDF2 no disponible'
            
            # Análisis de strings del PDF (búsqueda de URLs, emails, etc.)
            result['extracted_data'] = self._extract_strings_from_pdf(file_path)
            
        except Exception as e:
            result['error'] = f'Error al analizar PDF: {str(e)}'
            print_error(result['error'])
        
        return result
    
    def _extract_strings_from_pdf(self, file_path: str) -> Dict:
        """Extrae strings interesantes del PDF"""
        extracted = {
            'urls': [],
            'emails': [],
            'ips': [],
            'phone_numbers': [],
            'software_mentions': []
        }
        
        try:
            # Leer archivo en modo binario y buscar strings
            with open(file_path, 'rb') as f:
                content = f.read()
                
                # Decodificar bytes a string, ignorando errores
                text = content.decode('utf-8', errors='ignore')
                
                # Buscar URLs
                url_pattern = r'https?://(?:[-\\w.])+(?:[:\\d]+)?(?:/(?:[\\w/_.])*(?:\\?(?:[\\w&=%.])*)?(?:#(?:[\\w.])*)?)?'
                extracted['urls'] = list(set(re.findall(url_pattern, text)))[:20]
                
                # Buscar emails
                email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\\.[a-zA-Z]{2,}'
                extracted['emails'] = list(set(re.findall(email_pattern, text)))[:20]
                
                # Buscar IPs
                ip_pattern = r'\\b(?:\\d{1,3}\\.){3}\\d{1,3}\\b'
                extracted['ips'] = list(set(re.findall(ip_pattern, text)))[:10]
                
                # Buscar software común
                software_keywords = [
                    'Microsoft Word', 'Adobe Acrobat', 'LibreOffice', 'OpenOffice',
                    'Microsoft Excel', 'PowerPoint', 'Photoshop', 'Illustrator',
                    'InDesign', 'AutoCAD', 'SketchUp', 'Revit', 'ArcGIS'
                ]
                
                for keyword in software_keywords:
                    if keyword.lower() in text.lower():
                        extracted['software_mentions'].append(keyword)
                
        except Exception as e:
            print_warning(f"No se pudieron extraer strings: {str(e)}")
        
        return extracted
    
    def analyze_office_document(self, file_path: str) -> Dict:
        """
        Analiza metadatos de documentos Office (Word, Excel, PowerPoint)
        
        Args:
            file_path: Ruta al documento
        
        Returns:
            Dict con metadatos
        """
        print_info(f"Analizando documento Office: {os.path.basename(file_path)}")
        
        result = {
            'file_path': file_path,
            'file_name': os.path.basename(file_path),
            'file_type': None,
            'metadata': {},
            'core_properties': {},
            'error': None,
            'extracted_data': {}
        }
        
        # Detectar tipo
        ext = os.path.splitext(file_path)[1].lower()
        result['file_type'] = self.supported_formats.get(ext, 'Unknown')
        
        try:
            # Intentar usar python-docx para Word
            if ext in ['.docx']:
                try:
                    from docx import Document
                    
                    doc = Document(file_path)
                    
                    # Propiedades del documento
                    if doc.core_properties:
                        cp = doc.core_properties
                        result['core_properties'] = {
                            'author': cp.author,
                            'category': cp.category,
                            'comments': cp.comments,
                            'content_status': cp.content_status,
                            'created': str(cp.created) if cp.created else None,
                            'identifier': cp.identifier,
                            'keywords': cp.keywords,
                            'language': cp.language,
                            'last_modified_by': cp.last_modified_by,
                            'last_printed': str(cp.last_printed) if cp.last_printed else None,
                            'modified': str(cp.modified) if cp.modified else None,
                            'revision': cp.revision,
                            'subject': cp.subject,
                            'title': cp.title,
                            'version': cp.version
                        }
                    
                    # Estadísticas del documento
                    result['document_stats'] = {
                        'paragraphs': len(doc.paragraphs),
                        'tables': len(doc.tables),
                        'sections': len(doc.sections)
                    }
                    
                    # Extraer texto para análisis
                    full_text = []
                    for para in doc.paragraphs[:50]:  # Limitar a 50 párrafos
                        full_text.append(para.text)
                    
                    text_content = '\\n'.join(full_text)
                    result['extracted_data'] = self._extract_patterns_from_text(text_content)
                    
                    print_success(f"Documento Word analizado: {result['document_stats']['paragraphs']} párrafos")
                    
                except ImportError:
                    print_warning("python-docx no instalado")
                    result['error'] = 'python-docx no disponible'
            
            # Intentar usar openpyxl para Excel
            elif ext in ['.xlsx']:
                try:
                    from openpyxl import load_workbook
                    
                    wb = load_workbook(file_path, read_only=True, data_only=True)
                    
                    # Propiedades
                    if wb.properties:
                        props = wb.properties
                        result['core_properties'] = {
                            'creator': props.creator,
                            'title': props.title,
                            'description': props.description,
                            'subject': props.subject,
                            'identifier': props.identifier,
                            'language': props.language,
                            'created': str(props.created) if props.created else None,
                            'modified': str(props.modified) if props.modified else None,
                            'lastModifiedBy': props.lastModifiedBy,
                            'category': props.category,
                            'contentStatus': props.contentStatus,
                            'version': props.version,
                            'revision': props.revision,
                            'keywords': props.keywords
                        }
                    
                    # Información de hojas
                    result['workbook_info'] = {
                        'sheet_names': wb.sheetnames,
                        'total_sheets': len(wb.sheetnames)
                    }
                    
                    # Contar filas en cada hoja (limitado)
                    sheet_stats = {}
                    for sheet_name in wb.sheetnames[:3]:  # Limitar a 3 hojas
                        ws = wb[sheet_name]
                        row_count = 0
                        for row in ws.iter_rows():
                            row_count += 1
                            if row_count >= 1000:  # Limitar conteo
                                break
                        sheet_stats[sheet_name] = {'approx_rows': row_count}
                    
                    result['workbook_info']['sheet_stats'] = sheet_stats
                    
                    print_success(f"Libro Excel analizado: {result['workbook_info']['total_sheets']} hojas")
                    
                except ImportError:
                    print_warning("openpyxl no instalado")
                    result['error'] = 'openpyxl no disponible'
            
            # Para archivos legacy (.doc, .xls, .ppt), usar olefile
            elif ext in ['.doc', '.xls', '.ppt']:
                try:
                    import olefile
                    
                    if olefile.isOleFile(file_path):
                        ole = olefile.OleFileIO(file_path)
                        
                        # Metadatos del archivo OLE
                        meta = ole.get_metadata()
                        if meta:
                            result['ole_metadata'] = {
                                'author': meta.author,
                                'code_page': meta.code_page,
                                'comments': meta.comments,
                                'company': meta.company,
                                'creating_application': meta.creating_application,
                                'last_saved_by': meta.last_saved_by,
                                'locale_id': meta.locale_id,
                                'manager': meta.manager,
                                'subject': meta.subject,
                                'title': meta.title
                            }
                        
                        # Listar streams
                        result['ole_streams'] = ole.listdir()
                        
                        ole.close()
                        
                        print_success(f"Archivo OLE analizado: {len(result['ole_streams'])} streams")
                    
                except ImportError:
                    print_warning("olefile no instalado")
                    result['error'] = 'olefile no disponible'
            
            # Análisis genérico para todos los tipos
            result['file_stats'] = {
                'size_bytes': os.path.getsize(file_path),
                'size_human': self._human_readable_size(os.path.getsize(file_path)),
                'created': datetime.fromtimestamp(os.path.getctime(file_path)).isoformat(),
                'modified': datetime.fromtimestamp(os.path.getmtime(file_path)).isoformat()
            }
            
        except Exception as e:
            result['error'] = f'Error al analizar documento: {str(e)}'
            print_error(result['error'])
        
        return result
    
    def _extract_patterns_from_text(self, text: str) -> Dict:
        """Extrae patrones de texto"""
        patterns = {
            'emails': [],
            'urls': [],
            'phone_numbers': [],
            'ips': [],
            'dates': [],
            'credit_cards': [],
            'ssn': []
        }
        
        # Emails
        email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\\.[a-zA-Z]{2,}'
        patterns['emails'] = list(set(re.findall(email_pattern, text)))[:10]
        
        # URLs
        url_pattern = r'https?://(?:[-\\w.])+(?:[:\\d]+)?(?:/(?:[\\w/_.])*)?'
        patterns['urls'] = list(set(re.findall(url_pattern, text)))[:10]
        
        # Teléfonos (formato internacional básico)
        phone_pattern = r'\\+?\\d{1,3}[-.\\s]?\\(?\\d{1,4}\\)?[-.\\s]?\\d{1,4}[-.\\s]?\\d{1,9}'
        patterns['phone_numbers'] = list(set(re.findall(phone_pattern, text)))[:10]
        
        # IPs
        ip_pattern = r'\\b(?:\\d{1,3}\\.){3}\\d{1,3}\\b'
        patterns['ips'] = list(set(re.findall(ip_pattern, text)))[:10]
        
        # Fechas (formatos comunes)
        date_patterns = [
            r'\\d{1,2}/\\d{1,2}/\\d{2,4}',  # MM/DD/YYYY
            r'\\d{1,2}-\\d{1,2}-\\d{2,4}',   # MM-DD-YYYY
            r'\\d{4}-\\d{2}-\\d{2}',          # YYYY-MM-DD
        ]
        for dp in date_patterns:
            patterns['dates'].extend(re.findall(dp, text))
        patterns['dates'] = list(set(patterns['dates']))[:10]
        
        return patterns
    
    def _human_readable_size(self, size_bytes: int) -> str:
        """Convierte bytes a formato legible"""
        for unit in ['B', 'KB', 'MB', 'GB', 'TB']:
            if size_bytes < 1024.0:
                return f"{size_bytes:.2f} {unit}"
            size_bytes /= 1024.0
        return f"{size_bytes:.2f} PB"
    
    def analyze_document(self, file_path: str) -> Dict:
        """
        Analiza cualquier tipo de documento soportado
        
        Args:
            file_path: Ruta al documento
        
        Returns:
            Dict con análisis completo
        """
        print_banner(f"ANÁLISIS DE DOCUMENTO: {os.path.basename(file_path)}")
        
        if not os.path.exists(file_path):
            print_error(f"Archivo no encontrado: {file_path}")
            return {'error': 'Archivo no encontrado'}
        
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            result = self.analyze_pdf(file_path)
        elif ext in ['.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx', '.odt', '.ods', '.odp', '.rtf']:
            result = self.analyze_office_document(file_path)
        else:
            print_error(f"Formato no soportado: {ext}")
            return {'error': f'Formato no soportado: {ext}'}
        
        # Mostrar resultados
        print(f"\n{Colors.BOLD}INFORMACIÓN DEL ARCHIVO:{Colors.ENDC}")
        print(f"  Nombre: {result['file_name']}")
        print(f"  Tipo: {result.get('file_type', 'Desconocido')}")
        
        if result.get('file_stats'):
            fs = result['file_stats']
            print(f"  Tamaño: {fs['size_human']} ({fs['size_bytes']} bytes)")
            print(f"  Creado: {fs['created']}")
            print(f"  Modificado: {fs['modified']}")
        
        # Metadatos
        if result.get('core_properties'):
            print(f"\n{Colors.BOLD}PROPIEDADES DEL DOCUMENTO:{Colors.ENDC}")
            for key, value in result['core_properties'].items():
                if value:
                    print(f"  {key}: {value}")
        
        if result.get('metadata'):
            print(f"\n{Colors.BOLD}METADATOS:{Colors.ENDC}")
            for key, value in result['metadata'].items():
                if value:
                    print(f"  {key}: {value}")
        
        # Datos extraídos
        if result.get('extracted_data'):
            ed = result['extracted_data']
            print(f"\n{Colors.BOLD}DATOS EXTRAÍDOS:{Colors.ENDC}")
            
            if ed.get('emails'):
                print(f"  Emails encontrados: {len(ed['emails'])}")
                for email in ed['emails'][:5]:
                    print(f"    • {email}")
            
            if ed.get('urls'):
                print(f"  URLs encontradas: {len(ed['urls'])}")
                for url in ed['urls'][:5]:
                    print(f"    • {url}")
            
            if ed.get('phone_numbers'):
                print(f"  Teléfonos encontrados: {len(ed['phone_numbers'])}")
        
        # Información específica del tipo
        if result.get('structure_info'):
            si = result['structure_info']
            print(f"\n{Colors.BOLD}ESTRUCTURA:{Colors.ENDC}")
            if 'num_pages' in si:
                print(f"  Páginas: {si['num_pages']}")
            if 'paragraphs' in si:
                print(f"  Párrafos: {si['paragraphs']}")
            if 'tables' in si:
                print(f"  Tablas: {si['tables']}")
        
        # Análisis de privacidad
        print(f"\n{Colors.BOLD}ANÁLISIS DE PRIVACIDAD:{Colors.ENDC}")
        privacy_issues = []
        
        if result.get('core_properties', {}).get('author'):
            privacy_issues.append("Autor del documento identificado")
        if result.get('core_properties', {}).get('last_modified_by'):
            privacy_issues.append("Último modificador identificado")
        if result.get('metadata', {}).get('creator'):
            privacy_issues.append("Software creador identificado")
        if result.get('extracted_data', {}).get('emails'):
            privacy_issues.append(f"Emails expuestos en el documento ({len(result['extracted_data']['emails'])})")
        
        if privacy_issues:
            print_warning("Información potencialmente sensible:")
            for issue in privacy_issues:
                print_warning(f"  • {issue}")
        else:
            print_success("No se detectaron metadatos personales obvios")
        
        self.results = result
        return result
    
    def batch_analyze(self, directory: str, recursive: bool = False) -> List[Dict]:
        """
        Analiza múltiples documentos en un directorio
        
        Args:
            directory: Directorio a procesar
            recursive: Si se debe procesar recursivamente
        
        Returns:
            Lista de resultados
        """
        print_banner(f"ANÁLISIS BATCH: {directory}")
        
        results = []
        files_to_process = []
        
        # Recopilar archivos
        if recursive:
            for root, dirs, files in os.walk(directory):
                for file in files:
                    ext = os.path.splitext(file)[1].lower()
                    if ext in self.supported_formats:
                        files_to_process.append(os.path.join(root, file))
        else:
            for file in os.listdir(directory):
                ext = os.path.splitext(file)[1].lower()
                if ext in self.supported_formats:
                    files_to_process.append(os.path.join(directory, file))
        
        print_info(f"Encontrados {len(files_to_process)} documentos")
        
        # Procesar cada archivo
        for i, file_path in enumerate(files_to_process, 1):
            print(f"\n{Colors.BOLD}[{i}/{len(files_to_process)}]{Colors.ENDC} {os.path.basename(file_path)}")
            result = self.analyze_document(file_path)
            results.append(result)
        
        # Resumen
        print(f"\n{Colors.BOLD}RESUMEN BATCH:{Colors.ENDC}")
        print(f"  Total procesados: {len(results)}")
        
        # Contar por tipo
        by_type = {}
        for r in results:
            ftype = r.get('file_type', 'Unknown')
            by_type[ftype] = by_type.get(ftype, 0) + 1
        
        for ftype, count in by_type.items():
            print(f"  {ftype}: {count}")
        
        return results
    
    def save_results(self, filename: str = None):
        """Guarda los resultados en JSON"""
        if not filename:
            filename = f"document_analysis_{self.results.get('file_name', 'unknown')}"
        return save_to_json(self.results, filename)


def main():
    """Función principal para uso en línea de comandos"""
    import argparse
    
    parser = argparse.ArgumentParser(
        description='Analizador de Metadatos de Documentos OSINT',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ejemplos:
  python document_analyzer.py -f documento.pdf
  python document_analyzer.py -f informe.docx --save
  python document_analyzer.py --batch ./documentos/ --recursive
        """
    )
    
    parser.add_argument('-f', '--file', help='Ruta al documento a analizar')
    parser.add_argument('--batch', help='Directorio con documentos a procesar')
    parser.add_argument('-r', '--recursive', action='store_true', help='Procesar recursivamente')
    parser.add_argument('--save', action='store_true', help='Guardar resultados en JSON')
    
    args = parser.parse_args()
    
    analyzer = DocumentAnalyzer()
    
    if args.file:
        results = analyzer.analyze_document(args.file)
        if args.save:
            analyzer.save_results()
    
    elif args.batch:
        results = analyzer.batch_analyze(args.batch, args.recursive)
        if args.save:
            save_to_json(results, f"batch_document_analysis")
    else:
        parser.print_help()


if __name__ == '__main__':
    main()
